# -*- coding: utf-8 -*-
"""把 _build/src/*.md（自定义标记）同时渲染为 论文3.docx 与 论文3.tex。

标记：
  @@H1 / @@H2 / @@H3 标题
  @@P   正文段落（$...$ 为行内公式）
  @@EQ  latex | (1)     行间公式，右侧编号
  @@TBL ... @@ENDTBL     表格
        caption: 表1 xxx
        cols: a | b | c
        row:  ... | ... | ...
        note: 注：xxx
  @@FIG  文件名 | 图1 xxx
  @@PB   分页
  @@ABS  摘要标题行（居中加粗）
"""
import os, re, sys, shutil, glob
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml.ns import qn, nsmap
from docx.oxml import OxmlElement
from lxml import etree
import latex2mathml.converter as l2m

HERE = os.path.dirname(os.path.abspath(__file__))
SRC = os.path.join(HERE, 'src')
OUT = os.path.dirname(HERE)                      # hzt/
FIGDIR = os.path.join(HERE, 'fig')

XSL = r"C:\Program Files\Microsoft Office\root\Office16\MML2OMML.XSL"
_transform = etree.XSLT(etree.parse(XSL))
NS_M = "http://schemas.openxmlformats.org/officeDocument/2006/math"

WARN = []


def latex_to_omml(latex):
    # latex2mathml 的 aligned 支持不完整，Word转换用等价的单列表格分行。
    word_latex = latex
    if r'\begin{aligned}' in word_latex:
        word_latex = word_latex.replace(r'\begin{aligned}', r'\begin{array}{l}')
        word_latex = word_latex.replace(r'\end{aligned}', r'\end{array}')
        word_latex = word_latex.replace('&J', r'\qquad J').replace('&', '')
    mathml = l2m.convert(word_latex)
    dom = etree.fromstring(mathml)
    if dom.xpath('//*[local-name()="merror"]'):
        raise ValueError('MathML conversion contains merror: '+latex)
    result = _transform(dom).getroot()
    for mathrun in result.xpath('.//m:r', namespaces={'m': NS_M}):
        props = mathrun.find(qn('w:rPr'))
        if props is None:
            props = OxmlElement('w:rPr')
            mathrun.insert(0, props)
        fonts = OxmlElement('w:rFonts')
        fonts.set(qn('w:ascii'), 'Cambria Math')
        fonts.set(qn('w:hAnsi'), 'Cambria Math')
        props.append(fonts)
    return result


# ---------------------------------------------------------------- 解析
def parse():
    blocks = []
    buf = []
    for fn in sorted(glob.glob(os.path.join(SRC, '*.md'))):
        with open(fn, encoding='utf-8') as f:
            buf.append(f.read())
    text = '\n'.join(buf)
    # @@INCLUDE 展开：从 src/_tables/ 内联表格片段
    tbdir = os.path.join(SRC, '_tables')
    def _inc(m):
        fp = os.path.join(tbdir, m.group(1).strip())
        with open(fp, encoding='utf-8') as fh:
            return fh.read().rstrip('\n')
    text = re.sub(r'^@@INCLUDE\s+(\S+)\s*$', _inc, text, flags=re.M)
    lines = text.split('\n')
    i = 0
    while i < len(lines):
        ln = lines[i].rstrip()
        if ln.startswith('@@TBL'):
            i += 1
            cap, cols, rows, note = '', [], [], ''
            while i < len(lines) and not lines[i].startswith('@@ENDTBL'):
                s = lines[i].strip()
                if s.startswith('caption:'):
                    cap = s[8:].strip()
                elif s.startswith('note:'):
                    note = s[5:].strip()
                elif s.startswith('cols:'):
                    cols = [c.strip() for c in s[5:].split('|')]
                elif s.startswith('row:'):
                    rows.append([c.strip() for c in s[4:].split('|')])
                i += 1
            i += 1
            blocks.append(('tbl', dict(caption=cap, cols=cols, rows=rows, note=note)))
            continue
        m = re.match(r'@@(H1|H2|H3|EQ|FIG|PB|ABS|P)\b\s*(.*)$', ln)
        if m:
            tag, rest = m.group(1), m.group(2).strip()
            if tag == 'EQ':
                if '|' in rest:
                    tex, num = rest.rsplit('|', 1)
                else:
                    tex, num = rest, ''
                blocks.append(('eq', tex.strip(), num.strip()))
            elif tag == 'FIG':
                f, c = [x.strip() for x in rest.split('|', 1)]
                blocks.append(('fig', f, c))
            elif tag == 'PB':
                blocks.append(('pb',))
            elif tag == 'ABS':
                blocks.append(('abs', rest))
            else:
                blocks.append((tag.lower(), rest))
        elif ln.strip():
            blocks.append(('p', ln.strip()))
        i += 1
    # 两组四列的宽表分别呈现，保留全部数值及原总表号。
    expanded = []
    for block in blocks:
        if block[0] == 'tbl' and len(block[1]['cols']) == 9:
            t = block[1]
            for label, indices in [('a', [0,1,2,3,4]), ('b', [0,5,6,7,8])]:
                copy = dict(t)
                copy['caption'] = re.sub(r'^(表\d+)', r'\1('+label+')', t['caption'])
                copy['cols'] = [t['cols'][i] for i in indices]
                copy['rows'] = [[row[i] for i in indices] for row in t['rows']]
                expanded.append(('tbl', copy))
        else:
            expanded.append(block)
    return expanded


INLINE = re.compile(r'\$(.+?)\$', re.S)


# ---------------------------------------------------------------- Word
def set_font(run, cn='宋体', en='Times New Roman', size=12, bold=False):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = en
    run._element.rPr.rFonts.set(qn('w:eastAsia'), cn)


BOLD = re.compile(r'\*\*(.+?)\*\*', re.S)
CITE = re.compile(r'\[\d+(?:[,，、-]\d+)*\]')


def text_runs(p, seg, size, cn, bold):
    """正文文献号用上角标，参考文献条目号保持基线。"""
    if p.style.name == 'Bibliography':
        r = p.add_run(seg)
        set_font(r, cn=cn, size=size, bold=bold)
        return
    pos = 0
    for m in CITE.finditer(seg):
        if m.start() > pos:
            r = p.add_run(seg[pos:m.start()])
            set_font(r, cn=cn, size=size, bold=bold)
        r = p.add_run(m.group())
        set_font(r, cn=cn, size=size, bold=bold)
        r.font.superscript = True
        pos = m.end()
    if pos < len(seg):
        r = p.add_run(seg[pos:])
        set_font(r, cn=cn, size=size, bold=bold)


def add_runs(p, text, size=12, cn='宋体', bold=False):
    """把含 $...$ 与 **粗体** 的文本写进段落，公式走 OMML。"""
    pos = 0
    for m in INLINE.finditer(text):
        if m.start() > pos:
            _plain(p, text[pos:m.start()], size, cn, bold)
        tex = m.group(1)
        try:
            p._p.append(latex_to_omml(tex))
        except Exception as e:
            WARN.append('行内公式失败: %s (%s)' % (tex, e))
            r = p.add_run(tex)
            set_font(r, cn=cn, size=size, bold=bold)
        pos = m.end()
    if pos < len(text):
        _plain(p, text[pos:], size, cn, bold)


def _plain(p, seg, size, cn, bold):
    """写不含行内公式的片段，处理 **粗体**。"""
    pos = 0
    for m in BOLD.finditer(seg):
        if m.start() > pos:
            text_runs(p, seg[pos:m.start()], size, cn, bold)
        text_runs(p, m.group(1), size, cn, True)
        pos = m.end()
    if pos < len(seg):
        text_runs(p, seg[pos:], size, cn, bold)


def body_para(doc, text, size=12, indent=True):
    p = doc.add_paragraph()
    reference = bool(re.match(r'^\[\d+\]\s', text))
    if reference:
        p.style = 'Bibliography'
    pf = p.paragraph_format
    pf.line_spacing = 1.2
    pf.space_after = Pt(0)
    pf.widow_control = True
    p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if reference:
        pf.left_indent = Cm(.65)
        pf.first_line_indent = Cm(-.65)
        pf.line_spacing = 1.1
        add_runs(p, text, size=10.5)
        return p
    if indent:
        pf.first_line_indent = Pt(size * 2)
    add_runs(p, text, size=size)
    return p


def build_docx(blocks, path, figmap):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.5)
    sec.top_margin = sec.bottom_margin = Cm(2.5)

    st = doc.styles['Normal']
    st.font.name = 'Times New Roman'
    st.font.size = Pt(12)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), '宋体')
    if 'Bibliography' not in doc.styles:
        doc.styles.add_style('Bibliography', WD_STYLE_TYPE.PARAGRAPH)
    for name, size in [('Heading 1',14),('Heading 2',12),('Heading 3',12)]:
        style = doc.styles[name]
        style.font.name = 'Times New Roman'
        style.font.size = Pt(size)
        style.font.bold = True
        style.element.get_or_add_rPr().get_or_add_rFonts().set(qn('w:eastAsia'), '黑体')
        style.paragraph_format.keep_with_next = True
    footer = sec.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    field = OxmlElement('w:fldSimple')
    field.set(qn('w:instr'), 'PAGE')
    footer._p.append(field)

    eqno = 0
    for b in blocks:
        k = b[0]
        if k == 'h1':
            p = doc.add_paragraph(style='Heading 1')
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(4)
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if b[1] == '参考文献':
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.page_break_before = True
            r = p.add_run(b[1])
            set_font(r, cn='黑体', size=14, bold=True)
        elif k == 'h2':
            p = doc.add_paragraph(style='Heading 2')
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(b[1])
            set_font(r, cn='黑体', size=12, bold=True)
        elif k == 'h3':
            p = doc.add_paragraph(style='Heading 3')
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(3)
            r = p.add_run(b[1])
            set_font(r, cn='黑体', size=12, bold=True)
        elif k == 'abs':
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_before = Pt(6)
            p.paragraph_format.space_after = Pt(6)
            r = p.add_run(b[1])
            set_font(r, cn='黑体', size=14 if b[1].replace(' ','')=='摘要' else 16, bold=True)
        elif k == 'p':
            body_para(doc, b[1])
        elif k == 'eq':
            tex, num = b[1], b[2]
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.first_line_indent = Pt(0)
            p.paragraph_format.left_indent = Pt(0)
            p.paragraph_format.right_indent = Pt(0)
            p.paragraph_format.keep_together = True
            p.paragraph_format.line_spacing = 1.15
            p.paragraph_format.space_before = Pt(3)
            p.paragraph_format.space_after = Pt(3)
            # 制表位：公式居中，编号右对齐
            pPr = p._p.get_or_add_pPr()
            tabs = OxmlElement('w:tabs')
            for pos, al in ((Cm(8.0), 'center'), (Cm(16.0), 'right')):
                t = OxmlElement('w:tab')
                t.set(qn('w:val'), al)
                t.set(qn('w:pos'), str(int(pos.twips)))
                tabs.append(t)
            pPr.append(tabs)
            r = p.add_run('\t')
            set_font(r)
            try:
                p._p.append(latex_to_omml(tex))
            except Exception as e:
                WARN.append('行间公式失败: %s (%s)' % (tex, e))
                r = p.add_run(tex)
                set_font(r)
            if num:
                r = p.add_run('\t' + num)
                set_font(r)
        elif k == 'tbl':
            t = b[1]
            if t['caption']:
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(8)
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.keep_with_next = True
                r = p.add_run(t['caption'])
                set_font(r, cn='黑体', size=10.5, bold=True)
            ncol = len(t['cols'])
            tab = doc.add_table(rows=0, cols=ncol)
            tab.style = 'Normal Table'
            tab.alignment = WD_TABLE_ALIGNMENT.CENTER
            tab.autofit = False
            fractions = ([.20,.60,.20] if ncol == 3 else [.36,.16,.16,.16,.16] if ncol == 5 else [1/ncol]*ncol)
            for col, fraction in zip(tab.columns, fractions):
                col.width = Cm(16*fraction)
            borders = OxmlElement('w:tblBorders')
            for edge in ['top','left','bottom','right','insideH','insideV']:
                el = OxmlElement('w:'+edge)
                el.set(qn('w:val'), 'single' if edge in ['top','bottom'] else 'nil')
                el.set(qn('w:sz'), '8')
                borders.append(el)
            tab._tbl.tblPr.append(borders)
            for row in [t['cols']] + t['rows']:
                cells = tab.add_row().cells
                trPr = cells[0]._tc.getparent().get_or_add_trPr()
                trPr.append(OxmlElement('w:cantSplit'))
                if len(tab.rows) == 1:
                    trPr.append(OxmlElement('w:tblHeader'))
                for ci, txt in enumerate(row[:ncol]):
                    cells[ci].width = Cm(16*fractions[ci])
                    if len(tab.rows) == 1:
                        cb = OxmlElement('w:tcBorders')
                        bottom = OxmlElement('w:bottom')
                        bottom.set(qn('w:val'), 'single')
                        bottom.set(qn('w:sz'), '4')
                        cb.append(bottom)
                        cells[ci]._tc.get_or_add_tcPr().append(cb)
                    cp = cells[ci].paragraphs[0]
                    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
                    cp.paragraph_format.line_spacing = 1.0
                    cp.paragraph_format.space_after = Pt(0)
                    add_runs(cp, txt, size=10 if ncol==3 else 9.5, bold=len(tab.rows)==1)
            if t['note']:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.line_spacing = 1.0
                add_runs(p, t['note'], size=9)
        elif k == 'fig':
            f, cap = b[1], b[2]
            src = figmap.get(f)
            if src and os.path.exists(src):
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p.paragraph_format.space_before = Pt(6)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.keep_with_next = True
                shape = p.add_run().add_picture(src, width=Cm(14.5))
                if shape.height > Cm(15):
                    ratio = Cm(15)/shape.height
                    shape.width = int(shape.width*ratio)
                    shape.height = Cm(15)
                shape._inline.docPr.set('descr', cap)
            else:
                WARN.append('缺图: %s' % f)
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(8)
            r = p.add_run(cap)
            set_font(r, size=10.5)
        elif k == 'pb':
            doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)
    doc.save(path)
    print('wrote', path)


# ---------------------------------------------------------------- LaTeX
TEX_PRE = r"""\documentclass[12pt,a4paper]{ctexart}
\usepackage{amsmath,amssymb,bm}
\usepackage{graphicx}
\graphicspath{{figures/}}
\usepackage{booktabs}
\usepackage{longtable}
\usepackage{array}
\usepackage{geometry}
\usepackage{caption}
\usepackage{float}
\usepackage{needspace}
\usepackage[hidelinks]{hyperref}
\geometry{left=2.5cm,right=2.5cm,top=2.5cm,bottom=2.5cm}
\captionsetup[figure]{labelsep=space,font=small}
\captionsetup[table]{labelsep=space,font=small}
\newcolumntype{C}[1]{>{\centering\arraybackslash}p{#1}}
\setlength{\parindent}{2em}
\linespread{1.08}
\renewcommand{\arraystretch}{1.2}
\title{论文3}
\date{}
\setCJKfamilyfont{hei}[AutoFakeBold=2]{SimHei}
\ctexset{section={format=\heiti\bfseries\zihao{4},beforeskip=10pt,afterskip=6pt},subsection={format=\heiti\bfseries\zihao{-4},beforeskip=8pt,afterskip=4pt},subsubsection={format=\heiti\bfseries\zihao{-4},beforeskip=6pt,afterskip=3pt}}
\begin{document}
\setlength{\abovedisplayskip}{6pt}
\setlength{\belowdisplayskip}{6pt}
"""

TEX_POST = "\n\\end{document}\n"


SPECIAL = {'%': r'\%', '&': r'\&', '#': r'\#', '_': r'\_',
           '~': r'\textasciitilde{}', '^': r'\textasciicircum{}',
           '∶': '$:$', '⁻': '$^{-}$', '⁺': '$^{+}$',
           '≤': r'$\le$', '≥': r'$\ge$', '×': r'$\times$',
           '≈': r'$\approx$', '±': r'$\pm$',
           '“': '``', '”': "''"}


def _esc_plain(seg):
    """转义非公式片段中的 LaTeX 特殊字符，并把 **粗体** 转成 \\textbf。"""
    out = []
    pos = 0
    for m in BOLD.finditer(seg):
        out.append(_esc_raw(seg[pos:m.start()]))
        out.append('\\textbf{%s}' % _esc_raw(m.group(1)))
        pos = m.end()
    out.append(_esc_raw(seg[pos:]))
    return ''.join(out)


def _esc_raw(s):
    out, pos = [], 0
    for m in CITE.finditer(s):
        out.append(''.join(SPECIAL.get(ch,ch) for ch in s[pos:m.start()]))
        out.append(r'\textsuperscript{'+m.group()+ '}')
        pos = m.end()
    out.append(''.join(SPECIAL.get(ch,ch) for ch in s[pos:]))
    return ''.join(out)


def esc(s):
    """把 $...$ 之外的部分转义，$...$ 内保持原样。"""
    if re.match(r'^\[\d+\]\s', s):
        return re.sub(r'https?://[^\s]+',lambda m:r'\url{'+m.group().rstrip('.')+'}'+('.' if m.group().endswith('.') else ''),''.join(SPECIAL.get(ch,ch) for ch in s))
    out, pos = [], 0
    for m in INLINE.finditer(s):
        out.append(_esc_plain(s[pos:m.start()]))
        out.append('$%s$' % m.group(1))
        pos = m.end()
    out.append(_esc_plain(s[pos:]))
    return ''.join(out)


def build_tex(blocks, path, figmap):
    out = [TEX_PRE]
    for b in blocks:
        k = b[0]
        if k == 'h1':
            if b[1] == '参考文献':
                out.append('\n\\clearpage\\begin{center}\\heiti\\bfseries\\zihao{4}参考文献\\end{center}\n')
            else:
                out.append('\n\\needspace{4\\baselineskip}\\section*{%s}\n' % esc(b[1]))
        elif k == 'h2':
            out.append('\n\\needspace{3\\baselineskip}\\subsection*{%s}\n' % esc(b[1]))
        elif k == 'h3':
            out.append('\n\\subsubsection*{%s}\n' % esc(b[1]))
        elif k == 'abs':
            out.append('\n\\begin{center}\\Large\\bfseries %s\\end{center}\n' % esc(b[1]))
        elif k == 'p':
            if re.match(r'^\[\d+\]\s',b[1]):
                out.append('\n{\\footnotesize\\noindent\\hangindent=0.65cm\\hangafter=1 %s\\par}\n' % esc(b[1]))
            else:
                out.append('\n%s\n' % esc(b[1]))
        elif k == 'eq':
            tex, num = b[1], b[2]
            if num:
                m = re.match(r'^\((\d+)\)$', num)
                if m:
                    out.append('\n\\begin{equation}\n%s\n\\end{equation}\n' % tex)
                    continue
            out.append('\n\\[ %s \\]\n' % tex)
        elif k == 'tbl':
            t = b[1]
            ncol = len(t['cols'])
            fractions = ([.20,.60,.20] if ncol == 3 else [.36,.16,.16,.16,.16] if ncol == 5 else [1/ncol]*ncol)
            spec = ''.join('C{\\dimexpr%.6f\\textwidth-2\\tabcolsep\\relax}' % x for x in fractions)
            long = len(t['rows']) >= 18
            out.append('\n{\\fontsize{9.5}{11.5}\\selectfont\n' if long else '\n\\begin{table}[H]\n\\centering\\fontsize{9.5}{11.5}\\selectfont\n')
            if long:
                out.append('\\begin{longtable}{%s}\n' % spec)
            if t['caption']:
                out.append('\\caption*{%s}%s\n' % (esc(t['caption']), '\\\\' if long else ''))
            if not long:
                out.append('\\begin{tabular}{%s}\n' % spec)
            out.append('\\toprule\n')
            header = ' & '.join('\\textbf{%s}' % esc(c) for c in t['cols']) + ' \\\\' + '\n\\midrule\n'
            out.append(header)
            if long:
                out.append('\\endfirsthead\n\\toprule\n'+header+'\\endhead\n')
            for row in t['rows']:
                out.append(' & '.join(esc(c) for c in (row+['']*ncol)[:ncol]) + ' \\\\' + '\n')
            out.append('\\bottomrule\n\\end{%s}\n' % ('longtable' if long else 'tabular'))
            if t['note']:
                out.append('{\\footnotesize\\par\\vspace{2pt}%s\\par}\n' % esc(t['note']))
            out.append('}\n' if long else '\\end{table}\n')
        elif k == 'fig':
            f, cap = b[1], b[2]
            out.append('\n\\begin{figure}[H]\n\\centering\n'
                       '\\includegraphics[width=0.90\\textwidth,height=0.55\\textheight,keepaspectratio]{%s}\n'
                       '\\caption*{%s}\n\\end{figure}\n' % (f, esc(cap)))
        elif k == 'pb':
            out.append('\n\\clearpage\n')
    out.append(TEX_POST)
    with open(path, 'w', encoding='utf-8') as f:
        f.write(''.join(out))
    print('wrote', path)


def main():
    blocks = parse()
    figmap = {}
    # 图件映射：src 中写的名字 -> 实际文件
    for pat, tgt in [
        (os.path.join(FIGDIR, '*'), None),
    ]:
        pass
    for f in glob.glob(os.path.join(FIGDIR, '*')):
        figmap[os.path.basename(f)] = f
    for f in glob.glob(os.path.join(os.path.dirname(HERE), '论文3_LaTeX', 'figures', '*')):
        figmap[os.path.basename(f)] = f
    build_tex(blocks, os.path.join(OUT, '论文3_LaTeX', '论文3.tex'), figmap)
    try:
        build_docx(blocks, os.path.join(OUT, '论文3.docx'), figmap)
    except PermissionError:
        print('!! 论文3.docx 被占用（可能在 Word 中打开），本轮未更新 docx；'
              '关闭 Word 后重跑 mk.py 即可。tex 已正常更新。')
    if WARN:
        print('--- WARNINGS ---')
        for w in WARN:
            print(' ', w)


if __name__ == '__main__':
    main()
