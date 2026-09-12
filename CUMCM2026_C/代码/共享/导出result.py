"""按 附件5 模板生成 结果/result/result*.xlsx。

本脚本**只读** 结果/ 下已经算好的明细（问题1 的完整结果 csv、问题2/3/4 的
逐十分钟策略与逐日汇总、问题3/4-3 的链式承诺路径），不重算任何决策、不调用
求解器。因此它可以在全年计算完成后随时重跑，产出与论文口径一致的提交表格。

填表口径与已冻结的问题3 exporter 保持一致：
  - 模板第 2..145 列共 144 个时段列，按**时段编号**与逐十分钟明细一一对应
    （表头文字按区间右端点命名，即第 t 列标签为右端点，本脚本不因标签偏移而移位）；
  - 第 146 列 = 全天购电量，第 147 列 = 全天购电费；
  - result3 / result4-3 多一张「调整购电量」，填三次调整的变动绝对值之和；
  - 「充放电量」按 0:00-4:00 … 20:00-24:00 六段汇总，并填当日 0:00、24:00 储电量；
  - 「紧急购电量」把逐十分钟紧急购电合并为连续区间。

产出：
    结果/result/result1.xlsx    问题1（约束A 固定端点口径）
    结果/result/result2.xlsx    问题2
    结果/result/result3.xlsx    问题3
    结果/result/result4-2.xlsx  问题4-2
    结果/result/result4-3.xlsx  问题4-3

用法：python 导出result.py
"""

from __future__ import annotations

import shutil
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd

HERE = Path(__file__).resolve().parent            # 代码/共享
PROJECT = HERE.parent.parent                      # 项目根
ATTACH = next((b / "C题" / "附件" for b in (PROJECT, *PROJECT.parents)
               if (b / "C题" / "附件" / "附件1.xlsx").exists()), None)
if ATTACH is None:
    raise FileNotFoundError("未找到 附件1.xlsx，请保持 C题/附件/ 目录完整")
TEMPLATE_DIR = ATTACH / "附件5"
RES = PROJECT / "结果"
OUT = RES / "result"
OUT.mkdir(parents=True, exist_ok=True)

N_SLOT = 144
SLOTS_PER_HOUR = 6
BLOCKS = [(0, 4), (4, 8), (8, 12), (12, 16), (16, 20), (20, 24)]
FORMAL_START = "2025-02-01"
N_FORMAL = 334
EMERGENCY_TOL = 1e-9
ROW_MODE = "sequential"          # 与 solve_q1_ab.py 一致


def _label(minutes: int) -> str:
    return f"{minutes // 60}:{minutes % 60:02d}"


def _merge_runs(slots: np.ndarray) -> list[tuple[int, int]]:
    if slots.size == 0:
        return []
    breaks = np.flatnonzero(np.diff(slots) > 1)
    starts = np.concatenate(([slots[0]], slots[breaks + 1]))
    ends = np.concatenate(([slots[breaks], [slots[-1]]]))
    return list(zip(starts.tolist(), ends.tolist()))


def _load(problem: str) -> tuple[pd.DataFrame, pd.DataFrame, np.ndarray | None]:
    folder = problem.replace("Q", "问题")
    d = RES / "问题4" / folder if problem.startswith("Q4") else RES / folder
    daily = pd.read_csv(d / f"{problem}_daily.csv", parse_dates=["date"])
    slots = pd.read_csv(d / f"{problem}_physical_10min.csv", parse_dates=["date"])
    daily = daily[daily.date >= FORMAL_START].sort_values("date").reset_index(drop=True)
    slots = slots[slots.date >= FORMAL_START].sort_values(["date", "physical_slot0"]).reset_index(drop=True)
    paths = None
    npz = d / f"{problem}_plans.npz"
    with np.load(npz) as z:
        dates = pd.DatetimeIndex(z["dates"])
        keep = dates >= FORMAL_START
        paths = z["paths"][keep]
    if len(daily) != N_FORMAL or len(slots) != N_FORMAL * N_SLOT or len(paths) != N_FORMAL:
        raise ValueError(f"{problem} 正式期行数异常：daily={len(daily)}, slots={len(slots)}, paths={len(paths)}")
    return daily, slots, paths


def _series(slots: pd.DataFrame, column: str) -> np.ndarray:
    return slots[column].to_numpy(float).reshape(N_FORMAL, N_SLOT)


def _fill_plan_sheets(wb, dates, plan: np.ndarray, adjust: np.ndarray | None,
                      planned_cost: np.ndarray, adjust_cost: np.ndarray | None) -> None:
    sheets = [("计划购电量", plan, planned_cost)]
    if adjust is not None:
        sheets.append(("调整购电量", adjust, adjust_cost))
    for sheet, series, cost in sheets:
        ws = wb[sheet]
        for i in range(len(dates)):
            row = i + 2
            ws.cell(row, 1).value = pd.Timestamp(dates[i]).strftime("%Y-%m-%d")
            for t in range(N_SLOT):
                ws.cell(row, 2 + t).value = round(float(series[i, t]), 4)
            ws.cell(row, 2 + N_SLOT).value = round(float(series[i].sum()), 4)
            ws.cell(row, 3 + N_SLOT).value = round(float(cost[i]), 2)


def _fill_storage_sheet(wb, dates, charge: np.ndarray, discharge: np.ndarray,
                        soc_start: np.ndarray, soc_end: np.ndarray) -> None:
    ws = wb["充放电量"]
    for row in range(2, ws.max_row + 1):
        for col in range(1, 7):
            ws.cell(row, col).value = None
    row = 2
    for i in range(len(dates)):
        for b, (h0, h1) in enumerate(BLOCKS):
            sl = slice(h0 * SLOTS_PER_HOUR, h1 * SLOTS_PER_HOUR)
            ws.cell(row, 1).value = pd.Timestamp(dates[i]).strftime("%Y-%m-%d") if b == 0 else None
            ws.cell(row, 2).value = f"{h0}:00-{h1}:00"
            ws.cell(row, 3).value = round(float(charge[i, sl].sum()), 4)
            ws.cell(row, 4).value = round(float(discharge[i, sl].sum()), 4)
            if b == 0:
                ws.cell(row, 5).value = "0:00"
                ws.cell(row, 6).value = round(float(soc_start[i]), 4)
            elif b == 1:
                ws.cell(row, 5).value = "24:00"
                ws.cell(row, 6).value = round(float(soc_end[i]), 4)
            row += 1


def _fill_emergency_sheet(wb, dates, emergency: np.ndarray) -> None:
    ws = wb["紧急购电量"]
    for row in range(ws.max_row, 1, -1):
        for col in range(1, 4):
            ws.cell(row, col).value = None
    row = 2
    for i in range(len(dates)):
        runs = _merge_runs(np.flatnonzero(emergency[i] > EMERGENCY_TOL) + 1)
        for j, (s, e) in enumerate(runs):
            ws.cell(row, 1).value = pd.Timestamp(dates[i]).strftime("%Y-%m-%d") if j == 0 else None
            ws.cell(row, 2).value = f"{_label((s - 1) * 10)}-{_label(e * 10)}"
            ws.cell(row, 3).value = round(float(emergency[i, s - 1:e].sum()), 4)
            row += 1


def export_chain(problem: str, tag: str, staged: bool) -> Path:
    daily, slots, paths = _load(problem)
    dates = daily["date"].to_numpy()
    dest = OUT / f"result{tag}.xlsx"
    shutil.copyfile(TEMPLATE_DIR / f"result{tag}.xlsx", dest)
    wb = openpyxl.load_workbook(dest)

    plan = paths[:, 0, :]
    if staged:
        c = paths
        adjust = (np.abs(c[:, 1] - c[:, 0]) + np.abs(c[:, 2] - c[:, 1])
                  + np.abs(c[:, 3] - c[:, 2]))
        adjust_cost = (daily["breach_cost_yuan"] + daily["overbuy_cost_yuan"]).to_numpy(float)
    else:
        adjust = None
        adjust_cost = None
    _fill_plan_sheets(wb, dates, plan, adjust,
                      daily["base_cost_yuan"].to_numpy(float), adjust_cost)
    _fill_storage_sheet(wb, dates, _series(slots, "charge"), _series(slots, "discharge"),
                        daily["start_soc_kwh"].to_numpy(float),
                        daily["end_soc_kwh"].to_numpy(float))
    _fill_emergency_sheet(wb, dates, _series(slots, "emergency"))
    wb.save(dest)
    print(f"  {dest.relative_to(PROJECT)}  ({len(dates)} 天)")
    return dest


def export_q1():
    src = RES / "问题1" / "约束A（固定）_完整结果.csv"
    if not src.exists():
        print(f"  跳过 result1.xlsx：缺 {src.relative_to(PROJECT)}")
        print("  先跑：python 代码/问题1/solve_q1_ab.py")
        return None
    df = pd.read_csv(src, encoding="utf-8-sig")
    if len(df) != N_SLOT + 1:
        raise ValueError(f"问题1 完整结果行数不是 {N_SLOT + 1}，实际 {len(df)}")
    slots = df[df["时段 t"].between(1, N_SLOT)].reset_index(drop=True)
    if len(slots) != N_SLOT:
        raise ValueError(f"问题1 时段行数不是 {N_SLOT}，实际 {len(slots)}")
    x = slots["x_t·Δt (kWh)"].to_numpy(float)
    u = slots["u_t (kW)"].to_numpy(float) / 6.0
    v = slots["v_t (kW)"].to_numpy(float) / 6.0
    energy = df["E_t (kWh)"].to_numpy(float)      # 含 0:00 初始状态，共 T+1 点
    order = [(j + 1) % N_SLOT for j in range(N_SLOT)] if ROW_MODE == "label" else list(range(N_SLOT))

    dest = OUT / "result1.xlsx"
    shutil.copyfile(TEMPLATE_DIR / "result1.xlsx", dest)
    wb = openpyxl.load_workbook(dest)
    ws1 = wb["计划购电量"]
    for j, t in enumerate(order):
        ws1.cell(row=2 + j, column=2).value = round(float(x[t]), 4)
    ws2 = wb["充放电量"]
    for k in range(6):
        sl = slice(k * 24, (k + 1) * 24)
        ws2.cell(row=2 + k, column=2).value = round(float(u[sl].sum()), 4)
        ws2.cell(row=2 + k, column=3).value = round(float(v[sl].sum()), 4)
    ws2.cell(row=2, column=5).value = round(float(energy[0]), 4)
    ws2.cell(row=3, column=5).value = round(float(energy[-1]), 4)
    wb.save(dest)
    print(f"  {dest.relative_to(PROJECT)}  (1 天)")
    return dest


def main() -> None:
    print("按 附件5 模板导出 结果/result/ ：")
    export_q1()
    export_chain("Q2", "2", staged=False)
    export_chain("Q3", "3", staged=True)
    export_chain("Q4-2", "4-2", staged=False)
    export_chain("Q4-3", "4-3", staged=True)


if __name__ == "__main__":
    main()
