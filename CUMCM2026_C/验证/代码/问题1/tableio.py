# -*- coding: utf-8 -*-
"""
检验用的 Excel 写表工具。

一个表文件 = 一个 xlsx，里面可以放若干"块"（块 = 小标题 + 表头 + 数据行）。
检验表都是给人看的，所以做了这些排版：表题居中加粗、表头浅蓝底、数据行
居中、列宽按内容自适应（中文字按双宽算）、首行冻结、表注跟在数据下方。

被 check_model.py / check_optimal.py / sensitivity.py 共用。
"""

import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "微软雅黑"
HEADER_FILL = PatternFill("solid", fgColor="DCE6F1")
HEADER_SIDE = Side(style="thin", color="9E9E9E")


def _text_width(s):
    """中文按两格宽估算列宽"""
    return sum(2 if ord(ch) > 127 else 1 for ch in str(s))


def write_table(path, title, blocks, notes=None):
    """
    path    : 输出 xlsx 路径
    title   : 表题（写在第一行，跨列居中）
    blocks  : 列表，每项形如
              {"subtitle": "小标题或 None",
               "header":   ["列1", "列2", ...],
               "rows":     [[v1, v2, ...], ...],
               "formats":  [["0.0000", None, ...], ...] 或 None}
    notes   : 表注列表，每条一行，写在数据下方
    """
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title[:28] if title else "表"

    ncol = max(len(b["header"]) for b in blocks)
    row = 1

    # ---- 表题
    ws.cell(row=row, column=1, value=title)
    ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
    c = ws.cell(row=row, column=1)
    c.font = Font(name=FONT, bold=True, size=13)
    c.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[row].height = 26
    row += 1

    widths = [10] * ncol          # 逐列记录内容宽度
    first_data_row = None

    for bi, b in enumerate(blocks):
        header = b["header"]
        rows = b["rows"]
        formats = b.get("formats") or [None] * len(rows)

        if b.get("subtitle"):
            row += 1
            ws.cell(row=row, column=1, value=b["subtitle"])
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
            sc = ws.cell(row=row, column=1)
            sc.font = Font(name=FONT, bold=True, size=11, color="333333")
            sc.alignment = Alignment(horizontal="left", vertical="center")

        row += 1
        for j, h in enumerate(header, start=1):
            cell = ws.cell(row=row, column=j, value=h)
            cell.font = Font(name=FONT, bold=True, size=10)
            cell.fill = HEADER_FILL
            cell.border = Border(bottom=HEADER_SIDE)
            cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
            widths[j - 1] = max(widths[j - 1], _text_width(h))
        ws.row_dimensions[row].height = 22

        for i, data in enumerate(rows):
            row += 1
            if first_data_row is None:
                first_data_row = row
            frow = formats[i] if i < len(formats) and formats[i] else []
            for j, v in enumerate(data, start=1):
                cell = ws.cell(row=row, column=j, value=v)
                cell.font = Font(name=FONT, size=10)
                cell.alignment = Alignment(
                    horizontal="left" if j == 1 else "center", vertical="center")
                if j - 1 < len(frow) and frow[j - 1] and isinstance(v, (int, float)):
                    cell.number_format = frow[j - 1]
                if v is not None:
                    widths[j - 1] = max(widths[j - 1], _text_width(v))

    # ---- 表注
    if notes:
        row += 1
        ws.cell(row=row, column=1, value="说明：").font = Font(name=FONT, bold=True, size=9)
        for n in notes:
            row += 1
            ws.cell(row=row, column=1, value=n)
            ws.merge_cells(start_row=row, start_column=1, end_row=row, end_column=ncol)
            nc = ws.cell(row=row, column=1)
            nc.font = Font(name=FONT, size=9, color="555555")
            nc.alignment = Alignment(horizontal="left", vertical="center", wrap_text=False)
            widths[0] = max(widths[0], _text_width(n))

    # ---- 列宽与冻结
    for j in range(ncol):
        ws.column_dimensions[get_column_letter(j + 1)].width = min(60, max(9, widths[j] + 3))
    if first_data_row:
        ws.freeze_panes = ws.cell(row=first_data_row, column=1)

    try:
        wb.save(path)
    except PermissionError:
        print("  写入失败：%s" % path)
        print("  该文件正被其他程序占用（多半是在 Excel 里开着），关掉后重跑一次即可。")
        return None
    print("  已写出：", path)
    return path
